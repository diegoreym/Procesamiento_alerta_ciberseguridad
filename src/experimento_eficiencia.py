import csv
import os
import platform
import statistics
import sys
import time
from pathlib import Path
from random import Random

from modelos import Alerta
from bst import ArbolBST
from avl import ArbolAVL
from heap_prioridad import HeapPrioridad
from datos import generar_alertas


CARPETA_RESULTADOS: Path = Path(__file__).resolve().parent.parent / "resultados"
ESCENARIOS: tuple[str, str] = ("creciente", "aleatorio")


class ExperimentoEficiencia:
    """Ejecuta las mediciones, guarda los CSV y genera los graficos."""

    def __init__(
        self,
        tamanios: tuple[int, ...] = (100, 500, 1000),
        semilla: int = 2026,
        repeticiones: int = 7,
        consultas: int = 50,
        carpeta_salida: str | Path | None = None,
    ) -> None:
        if not tamanios or any(type(n) is not int or n <= 0 for n in tamanios):
            raise ValueError("Los tamanios deben ser enteros positivos")
        if type(semilla) is not int:
            raise TypeError("La semilla debe ser un entero")
        if type(repeticiones) is not int or repeticiones < 1:
            raise ValueError("Las repeticiones deben ser un entero positivo")
        if type(consultas) is not int or not 1 <= consultas <= min(tamanios):
            raise ValueError("Las consultas deben ser un entero entre 1 y el menor tamanio")

        self.__tamanios: tuple[int, ...] = tuple(tamanios)
        self.__semilla: int = semilla
        self.__repeticiones: int = repeticiones
        self.__consultas: int = consultas
        self.__carpeta: Path = (
            Path(carpeta_salida) if carpeta_salida is not None else CARPETA_RESULTADOS
        )

    @property
    def tamanios(self) -> tuple[int, ...]:
        return self.__tamanios

    @property
    def semilla(self) -> int:
        return self.__semilla

    @property
    def repeticiones(self) -> int:
        return self.__repeticiones

    @property
    def consultas(self) -> int:
        return self.__consultas

    @property
    def carpeta_salida(self) -> Path:
        return self.__carpeta

    # ------------------------------------------------------------------
    # Preparacion de datos (siempre fuera del cronometro)
    # ------------------------------------------------------------------
    def generar_orden(self, n: int, escenario: str) -> list[Alerta]:
        """Alertas con ids 1..n en el orden de insercion del escenario."""
        alertas = generar_alertas(n, self.__semilla)
        if escenario == "creciente":
            return alertas
        if escenario == "aleatorio":
            Random(self.__semilla).shuffle(alertas)
            return alertas
        raise ValueError(f"Escenario desconocido: {escenario}")

    def seleccionar_consultas(self, n: int) -> list[int]:
        """Ids existentes a buscar; la misma lista se usa en BST y AVL."""
        return Random(self.__semilla).sample(range(1, n + 1), self.__consultas)

    @staticmethod
    def __verificar_consultas(arbol, ids: list[int], nombre: str) -> None:
        for id in ids:
            alerta = arbol.buscar(id)
            if alerta is None or alerta.id != id:
                raise RuntimeError(f"{nombre} no encontro el id {id}")

    @staticmethod
    def __cronometrar_busquedas(arbol, ids: list[int]) -> float:
        buscar = arbol.buscar
        inicio = time.perf_counter()
        for id in ids:
            buscar(id)
        return time.perf_counter() - inicio

    # ------------------------------------------------------------------
    # Arboles: altura y busqueda
    # ------------------------------------------------------------------
    def medir_arboles(self) -> tuple[list[dict], list[dict], list[dict]]:
        """Devuelve (filas de arboles.csv, mediciones crudas, ordenes usados)."""
        filas: list[dict] = []
        crudas: list[dict] = []
        ordenes: list[dict] = []
        operacion = f"buscar_lote_{self.__consultas}"

        for n in self.__tamanios:
            ids_consulta = self.seleccionar_consultas(n)

            for escenario in ESCENARIOS:
                orden = self.generar_orden(n, escenario)
                ordenes.append({
                    "n": n,
                    "escenario": escenario,
                    "semilla": self.__semilla,
                    "orden_ids": " ".join(str(a.id) for a in orden),
                })

                bst = ArbolBST()
                avl = ArbolAVL()
                for alerta in orden:
                    bst.insertar(alerta)
                    avl.insertar(alerta)

                if not avl.validar_balance():
                    raise RuntimeError(f"El AVL perdio el balance (n={n}, {escenario})")
                self.__verificar_consultas(bst, ids_consulta, "BST")
                self.__verificar_consultas(avl, ids_consulta, "AVL")

                altura_bst = bst.altura()
                altura_avl = avl.altura()

                # Una pasada de calentamiento sin registrar.
                self.__cronometrar_busquedas(bst, ids_consulta)
                self.__cronometrar_busquedas(avl, ids_consulta)

                tiempos: dict[str, list[float]] = {"BST": [], "AVL": []}
                for repeticion in range(1, self.__repeticiones + 1):
                    # Se alterna quien se mide primero para no favorecer a uno.
                    if repeticion % 2 == 1:
                        pares = (("BST", bst), ("AVL", avl))
                    else:
                        pares = (("AVL", avl), ("BST", bst))
                    for nombre, arbol in pares:
                        tiempo = self.__cronometrar_busquedas(arbol, ids_consulta)
                        tiempos[nombre].append(tiempo)
                        crudas.append({
                            "repeticion": repeticion,
                            "estructura": nombre,
                            "escenario": escenario,
                            "n": n,
                            "operacion": operacion,
                            "tiempo_s": tiempo,
                        })

                filas.append({
                    "n": n,
                    "escenario": escenario,
                    "semilla": self.__semilla,
                    "altura_bst": altura_bst,
                    "altura_avl": altura_avl,
                    "mediana_busqueda_bst_s": statistics.median(tiempos["BST"]),
                    "mediana_busqueda_avl_s": statistics.median(tiempos["AVL"]),
                })

        return filas, crudas, ordenes

    # ------------------------------------------------------------------
    # Heap: insertar n y extraer todas
    # ------------------------------------------------------------------
    @staticmethod
    def __verificar_orden_heap(alertas: list[Alerta]) -> None:
        heap = HeapPrioridad()
        for alerta in alertas:
            heap.insertar(alerta)
        obtenido = [heap.extraer().id for _ in range(len(alertas))]
        esperado = [
            a.id for a in sorted(alertas, key=lambda a: (-a.riesgo, a.instante, a.id))
        ]
        if obtenido != esperado:
            raise RuntimeError("El heap no extrajo las alertas en el orden esperado")

    def medir_heap(self) -> tuple[list[dict], list[dict]]:
        """Devuelve (filas de heap.csv, mediciones crudas)."""
        filas: list[dict] = []
        crudas: list[dict] = []

        for n in self.__tamanios:
            alertas = generar_alertas(n, self.__semilla)
            self.__verificar_orden_heap(alertas)

            insertar_s: list[float] = []
            extraer_s: list[float] = []

            for repeticion in range(1, self.__repeticiones + 1):
                heap = HeapPrioridad()

                inicio = time.perf_counter()
                for alerta in alertas:
                    heap.insertar(alerta)
                t_insertar = time.perf_counter() - inicio

                if heap.tamanio() != n:
                    raise RuntimeError("El heap no quedo con n alertas")

                inicio = time.perf_counter()
                for _ in range(n):
                    heap.extraer()
                t_extraer = time.perf_counter() - inicio

                if not heap.esta_vacia():
                    raise RuntimeError("El heap no quedo vacio tras extraer todo")

                insertar_s.append(t_insertar)
                extraer_s.append(t_extraer)
                for operacion, tiempo in (
                    ("insertar_n", t_insertar), ("extraer_n", t_extraer)
                ):
                    crudas.append({
                        "repeticion": repeticion,
                        "estructura": "HEAP",
                        "escenario": "generado",
                        "n": n,
                        "operacion": operacion,
                        "tiempo_s": tiempo,
                    })

            filas.append({
                "n": n,
                "semilla": self.__semilla,
                "mediana_insertar_s": statistics.median(insertar_s),
                "mediana_extraer_s": statistics.median(extraer_s),
            })

        return filas, crudas

    # ------------------------------------------------------------------
    # Salidas
    # ------------------------------------------------------------------
    @staticmethod
    def __escribir_csv(ruta: Path, campos: list[str], filas: list[dict]) -> None:
        with ruta.open("w", encoding="utf-8", newline="") as salida:
            escritor = csv.DictWriter(salida, fieldnames=campos)
            escritor.writeheader()
            for fila in filas:
                escritor.writerow({
                    campo: (f"{fila[campo]:.9f}" if isinstance(fila[campo], float)
                            else fila[campo])
                    for campo in campos
                })

    def escribir_entorno(self, ruta: Path) -> None:
        lineas = [
            f"Sistema operativo: {platform.platform()}",
            f"Procesador: {platform.processor() or 'no disponible'}",
            f"CPUs logicas: {os.cpu_count()}",
            "RAM: no disponible (completar a mano con los datos reales del equipo)",
            f"Python: {sys.version.split()[0]} ({platform.python_implementation()})",
            f"Interprete usado: {sys.executable}",
            f"Semilla: {self.__semilla}",
            f"Repeticiones: {self.__repeticiones}",
            f"Tamanios: {', '.join(str(n) for n in self.__tamanios)}",
            f"Consultas por lote: {self.__consultas}",
            "Medicion: time.perf_counter, mediana de las repeticiones",
        ]
        ruta.write_text("\n".join(lineas) + "\n", encoding="utf-8")

    @staticmethod
    def escribir_excel(
        ruta: Path, tablas: list[tuple[str, list[str], list[dict]]]
    ) -> None:
        """Escribe un .xlsx con una hoja por tabla, formateada como tabla de Excel.

        Los numeros se guardan como numeros (no como texto) para poder ordenarlos
        y calcular con ellos. Requiere openpyxl.
        """
        from openpyxl import Workbook
        from openpyxl.styles import Alignment
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo

        libro = Workbook()
        libro.remove(libro.active)

        for nombre, campos, filas in tablas:
            hoja = libro.create_sheet(nombre)
            hoja.append(campos)
            for fila in filas:
                hoja.append([fila[campo] for campo in campos])

            if filas:
                referencia = f"A1:{get_column_letter(len(campos))}{len(filas) + 1}"
                tabla = Table(displayName=f"tbl_{nombre}", ref=referencia)
                tabla.tableStyleInfo = TableStyleInfo(
                    name="TableStyleMedium2", showRowStripes=True
                )
                hoja.add_table(tabla)
            hoja.freeze_panes = "A2"

            for indice, campo in enumerate(campos, start=1):
                letra = get_column_letter(indice)
                textos = [campo] + [
                    f"{fila[campo]:.9f}" if isinstance(fila[campo], float)
                    else str(fila[campo])
                    for fila in filas
                ]
                hoja.column_dimensions[letra].width = min(
                    max(len(texto) for texto in textos) + 3, 60
                )
                for celda in hoja[letra][1:]:
                    if isinstance(celda.value, float):
                        celda.number_format = "0.000000000"
                    if isinstance(celda.value, (int, float)):
                        celda.alignment = Alignment(horizontal="center")

        libro.save(ruta)

    def graficar(self, filas_arboles: list[dict], filas_heap: list[dict]) -> list[Path]:
        """Genera alturas.png y tiempos.png. Requiere matplotlib."""
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt

        estilos = {
            ("BST", "creciente"): ("o", "-"),
            ("AVL", "creciente"): ("s", "-"),
            ("BST", "aleatorio"): ("o", "--"),
            ("AVL", "aleatorio"): ("s", "--"),
        }

        def serie(escenario: str, campo: str, factor: float = 1.0) -> list[float]:
            return [
                f[campo] * factor
                for f in sorted(filas_arboles, key=lambda f: f["n"])
                if f["escenario"] == escenario
            ]

        ns = sorted(self.__tamanios)

        # Alturas
        figura, eje = plt.subplots(figsize=(7, 4.5))
        for (nombre, escenario), (marca, linea) in estilos.items():
            campo = "altura_bst" if nombre == "BST" else "altura_avl"
            eje.plot(ns, serie(escenario, campo), marker=marca, linestyle=linea,
                     label=f"{nombre} {escenario}")
        eje.set_yscale("log")
        eje.set_xlabel("n (numero de alertas)")
        eje.set_ylabel("Altura en aristas (escala logaritmica)")
        eje.set_title("Altura de BST y AVL segun el orden de insercion")
        eje.set_xticks(ns)
        eje.grid(True, alpha=0.3)
        eje.legend()
        figura.tight_layout()
        ruta_alturas = self.__carpeta / "alturas.png"
        figura.savefig(ruta_alturas, dpi=150)
        plt.close(figura)

        # Tiempos
        figura, (eje1, eje2) = plt.subplots(1, 2, figsize=(12, 4.5))
        for (nombre, escenario), (marca, linea) in estilos.items():
            campo = ("mediana_busqueda_bst_s" if nombre == "BST"
                     else "mediana_busqueda_avl_s")
            eje1.plot(ns, serie(escenario, campo, 1000.0), marker=marca,
                      linestyle=linea, label=f"{nombre} {escenario}")
        eje1.set_xlabel("n (numero de alertas)")
        eje1.set_ylabel(f"Mediana (ms) de un lote de {self.__consultas} busquedas")
        eje1.set_title("Busqueda por id")
        eje1.set_xticks(ns)
        eje1.grid(True, alpha=0.3)
        eje1.legend()

        filas_h = sorted(filas_heap, key=lambda f: f["n"])
        eje2.plot(ns, [f["mediana_insertar_s"] * 1000.0 for f in filas_h],
                  marker="o", label="insertar n alertas")
        eje2.plot(ns, [f["mediana_extraer_s"] * 1000.0 for f in filas_h],
                  marker="s", label="extraer n alertas")
        eje2.set_xlabel("n (numero de alertas)")
        eje2.set_ylabel("Mediana (ms) de n operaciones")
        eje2.set_title("Heap de prioridad")
        eje2.set_xticks(ns)
        eje2.grid(True, alpha=0.3)
        eje2.legend()

        figura.tight_layout()
        ruta_tiempos = self.__carpeta / "tiempos.png"
        figura.savefig(ruta_tiempos, dpi=150)
        plt.close(figura)

        return [ruta_alturas, ruta_tiempos]

    def ejecutar(self) -> dict[str, Path]:
        """Corre todo el experimento y escribe los archivos en la carpeta de salida."""
        self.__carpeta.mkdir(parents=True, exist_ok=True)

        filas_arboles, crudas_arboles, ordenes = self.medir_arboles()
        filas_heap, crudas_heap = self.medir_heap()

        rutas: dict[str, Path] = {
            "arboles": self.__carpeta / "arboles.csv",
            "heap": self.__carpeta / "heap.csv",
            "crudas": self.__carpeta / "mediciones_crudas.csv",
            "ordenes": self.__carpeta / "ordenes_usados.csv",
            "entorno": self.__carpeta / "entorno.txt",
        }

        # (nombre, columnas, filas): se usa para los CSV y para las hojas de Excel.
        tablas: list[tuple[str, list[str], list[dict]]] = [
            ("arboles",
             ["n", "escenario", "semilla", "altura_bst", "altura_avl",
              "mediana_busqueda_bst_s", "mediana_busqueda_avl_s"],
             filas_arboles),
            ("heap",
             ["n", "semilla", "mediana_insertar_s", "mediana_extraer_s"],
             filas_heap),
            ("mediciones_crudas",
             ["repeticion", "estructura", "escenario", "n", "operacion", "tiempo_s"],
             crudas_arboles + crudas_heap),
            ("ordenes_usados",
             ["n", "escenario", "semilla", "orden_ids"],
             ordenes),
        ]
        for nombre, campos, filas in tablas:
            self.__escribir_csv(self.__carpeta / f"{nombre}.csv", campos, filas)
        self.escribir_entorno(rutas["entorno"])

        try:
            rutas["excel"] = self.__carpeta / "resultados.xlsx"
            self.escribir_excel(rutas["excel"], tablas)
        except ImportError:
            del rutas["excel"]
            print("Aviso: openpyxl no esta instalado; se omitio el Excel. "
                  "Instalelo con: pip install openpyxl")

        try:
            rutas_graficos = self.graficar(filas_arboles, filas_heap)
            rutas["alturas"], rutas["tiempos"] = rutas_graficos
        except ImportError:
            print("Aviso: matplotlib no esta instalado; se omitieron los graficos. "
                  "Instalelo con: pip install matplotlib")

        return rutas


if __name__ == "__main__":
    experimento = ExperimentoEficiencia()
    resultado = experimento.ejecutar()
    print("Experimento terminado. Archivos generados:")
    for nombre, ruta in resultado.items():
        print(f"  {nombre}: {ruta}")
