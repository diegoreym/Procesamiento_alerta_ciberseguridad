import csv
import math
import statistics
import sys
import tempfile
import unittest
from importlib.util import find_spec
from pathlib import Path

CARPETA_TESTS = Path(__file__).resolve().parent
for candidata in (CARPETA_TESTS.parent / "src", CARPETA_TESTS.parent):
    if (candidata / "experimento_eficiencia.py").is_file():
        sys.path.insert(0, str(candidata))
        break

from experimento_eficiencia import ExperimentoEficiencia  # noqa: E402


def leer_csv(ruta: Path) -> list[dict]:
    with ruta.open(encoding="utf-8", newline="") as entrada:
        return list(csv.DictReader(entrada))


class TestConfiguracion(unittest.TestCase):
    def test_valores_por_defecto(self):
        experimento = ExperimentoEficiencia()
        self.assertEqual(experimento.tamanios, (100, 500, 1000))
        self.assertEqual(experimento.semilla, 2026)
        self.assertEqual(experimento.consultas, 50)
        self.assertGreaterEqual(experimento.repeticiones, 5)

    def test_tamanios_invalidos(self):
        with self.assertRaises(ValueError):
            ExperimentoEficiencia(tamanios=())
        with self.assertRaises(ValueError):
            ExperimentoEficiencia(tamanios=(100, -5))

    def test_repeticiones_invalidas(self):
        with self.assertRaises(ValueError):
            ExperimentoEficiencia(repeticiones=0)

    def test_consultas_mayores_que_el_tamanio(self):
        with self.assertRaises(ValueError):
            ExperimentoEficiencia(tamanios=(20,), consultas=50)

    def test_semilla_no_entera(self):
        with self.assertRaises(TypeError):
            ExperimentoEficiencia(semilla="2026")


class TestPreparacionDatos(unittest.TestCase):
    def setUp(self):
        self.experimento = ExperimentoEficiencia()

    def test_orden_creciente(self):
        for n in (100, 500, 1000):
            ids = [a.id for a in self.experimento.generar_orden(n, "creciente")]
            self.assertEqual(ids, list(range(1, n + 1)))

    def test_orden_aleatorio_es_permutacion_de_los_mismos_ids(self):
        for n in (100, 500, 1000):
            ids = [a.id for a in self.experimento.generar_orden(n, "aleatorio")]
            self.assertEqual(len(ids), n)
            self.assertEqual(sorted(ids), list(range(1, n + 1)))
            self.assertNotEqual(ids, list(range(1, n + 1)))

    def test_orden_aleatorio_reproducible(self):
        primero = [a.id for a in self.experimento.generar_orden(500, "aleatorio")]
        segundo = [
            a.id
            for a in ExperimentoEficiencia().generar_orden(500, "aleatorio")
        ]
        self.assertEqual(primero, segundo)

    def test_otra_semilla_cambia_el_orden_aleatorio(self):
        base = [a.id for a in self.experimento.generar_orden(100, "aleatorio")]
        otro = [
            a.id
            for a in ExperimentoEficiencia(semilla=7).generar_orden(100, "aleatorio")
        ]
        self.assertNotEqual(base, otro)

    def test_escenario_desconocido(self):
        with self.assertRaises(ValueError):
            self.experimento.generar_orden(100, "inventado")

    def test_cincuenta_consultas_existentes_y_sin_repetir(self):
        for n in (100, 500, 1000):
            ids = self.experimento.seleccionar_consultas(n)
            self.assertEqual(len(ids), 50)
            self.assertEqual(len(set(ids)), 50)
            self.assertTrue(all(1 <= id <= n for id in ids))

    def test_consultas_reproducibles(self):
        self.assertEqual(
            self.experimento.seleccionar_consultas(1000),
            ExperimentoEficiencia().seleccionar_consultas(1000),
        )


class TestMediciones(unittest.TestCase):
    """Corre una vez un experimento pequeno y revisa sus filas."""

    TAMANIOS = (20, 50)
    REPETICIONES = 3
    CONSULTAS = 10

    @classmethod
    def setUpClass(cls):
        cls.experimento = ExperimentoEficiencia(
            tamanios=cls.TAMANIOS,
            repeticiones=cls.REPETICIONES,
            consultas=cls.CONSULTAS,
        )
        cls.filas_arboles, cls.crudas_arboles, cls.ordenes = (
            cls.experimento.medir_arboles()
        )
        cls.filas_heap, cls.crudas_heap = cls.experimento.medir_heap()

    def test_filas_de_arboles(self):
        self.assertEqual(len(self.filas_arboles), len(self.TAMANIOS) * 2)
        combinaciones = {(f["n"], f["escenario"]) for f in self.filas_arboles}
        esperadas = {(n, e) for n in self.TAMANIOS for e in ("creciente", "aleatorio")}
        self.assertEqual(combinaciones, esperadas)

    def test_alturas_en_aristas(self):
        for fila in self.filas_arboles:
            n = fila["n"]
            if fila["escenario"] == "creciente":
                self.assertEqual(fila["altura_bst"], n - 1)
                self.assertLess(fila["altura_avl"], fila["altura_bst"])
            self.assertLessEqual(fila["altura_avl"], 1.45 * math.log2(n + 2))
            self.assertGreaterEqual(fila["altura_avl"], math.floor(math.log2(n)))

    def test_mediciones_crudas_de_arboles(self):
        # 2 estructuras x 2 escenarios x repeticiones, por cada tamanio.
        esperadas = len(self.TAMANIOS) * 2 * 2 * self.REPETICIONES
        self.assertEqual(len(self.crudas_arboles), esperadas)
        self.assertTrue(all(f["tiempo_s"] >= 0 for f in self.crudas_arboles))
        self.assertEqual(
            {f["operacion"] for f in self.crudas_arboles},
            {f"buscar_lote_{self.CONSULTAS}"},
        )

    def test_mediana_de_arboles_sale_de_las_crudas(self):
        for fila in self.filas_arboles:
            for estructura, campo in (
                ("BST", "mediana_busqueda_bst_s"),
                ("AVL", "mediana_busqueda_avl_s"),
            ):
                tiempos = [
                    c["tiempo_s"]
                    for c in self.crudas_arboles
                    if c["n"] == fila["n"]
                    and c["escenario"] == fila["escenario"]
                    and c["estructura"] == estructura
                ]
                self.assertEqual(len(tiempos), self.REPETICIONES)
                self.assertEqual(fila[campo], statistics.median(tiempos))

    def test_orden_usado_guardado(self):
        self.assertEqual(len(self.ordenes), len(self.TAMANIOS) * 2)
        for registro in self.ordenes:
            ids = [int(x) for x in registro["orden_ids"].split()]
            self.assertEqual(sorted(ids), list(range(1, registro["n"] + 1)))
            if registro["escenario"] == "creciente":
                self.assertEqual(ids, sorted(ids))
            self.assertEqual(registro["semilla"], 2026)

    def test_filas_de_heap(self):
        self.assertEqual([f["n"] for f in self.filas_heap], list(self.TAMANIOS))
        for fila in self.filas_heap:
            self.assertGreaterEqual(fila["mediana_insertar_s"], 0)
            self.assertGreaterEqual(fila["mediana_extraer_s"], 0)

    def test_mediciones_crudas_del_heap(self):
        # insertar y extraer, por cada repeticion y tamanio.
        esperadas = len(self.TAMANIOS) * self.REPETICIONES * 2
        self.assertEqual(len(self.crudas_heap), esperadas)
        self.assertEqual(
            {f["operacion"] for f in self.crudas_heap}, {"insertar_n", "extraer_n"}
        )

    def test_mediana_del_heap_sale_de_las_crudas(self):
        for fila in self.filas_heap:
            for operacion, campo in (
                ("insertar_n", "mediana_insertar_s"),
                ("extraer_n", "mediana_extraer_s"),
            ):
                tiempos = [
                    c["tiempo_s"]
                    for c in self.crudas_heap
                    if c["n"] == fila["n"] and c["operacion"] == operacion
                ]
                self.assertEqual(len(tiempos), self.REPETICIONES)
                self.assertEqual(fila[campo], statistics.median(tiempos))


class TestEjecucionCompleta(unittest.TestCase):
    """Corre el protocolo con los tamanios reales (100, 500 y 1000)."""

    REPETICIONES = 2

    @classmethod
    def setUpClass(cls):
        cls.carpeta_temporal = tempfile.TemporaryDirectory(
            prefix="tmp_experimento_", dir=CARPETA_TESTS
        )
        cls.addClassCleanup(cls.carpeta_temporal.cleanup)
        cls.carpeta = Path(cls.carpeta_temporal.name) / "resultados"
        cls.experimento = ExperimentoEficiencia(
            repeticiones=cls.REPETICIONES, carpeta_salida=cls.carpeta
        )
        cls.rutas = cls.experimento.ejecutar()

    def test_crea_la_carpeta_y_los_archivos(self):
        self.assertTrue(self.carpeta.is_dir())
        for nombre in (
            "arboles.csv", "heap.csv", "mediciones_crudas.csv",
            "ordenes_usados.csv", "entorno.txt",
        ):
            self.assertTrue((self.carpeta / nombre).is_file(), nombre)

    def test_arboles_csv_tiene_seis_filas(self):
        filas = leer_csv(self.carpeta / "arboles.csv")
        self.assertEqual(len(filas), 6)
        self.assertEqual(
            list(filas[0]),
            ["n", "escenario", "semilla", "altura_bst", "altura_avl",
             "mediana_busqueda_bst_s", "mediana_busqueda_avl_s"],
        )
        creciente = {int(f["n"]): f for f in filas if f["escenario"] == "creciente"}
        for n in (100, 500, 1000):
            self.assertEqual(int(creciente[n]["altura_bst"]), n - 1)

    def test_heap_csv_tiene_tres_filas(self):
        filas = leer_csv(self.carpeta / "heap.csv")
        self.assertEqual([int(f["n"]) for f in filas], [100, 500, 1000])
        self.assertEqual(
            list(filas[0]),
            ["n", "semilla", "mediana_insertar_s", "mediana_extraer_s"],
        )

    def test_mediciones_crudas_csv(self):
        filas = leer_csv(self.carpeta / "mediciones_crudas.csv")
        arboles = 3 * 2 * 2 * self.REPETICIONES
        heap = 3 * 2 * self.REPETICIONES
        self.assertEqual(len(filas), arboles + heap)
        self.assertEqual({f["estructura"] for f in filas}, {"BST", "AVL", "HEAP"})

    def test_entorno_no_inventa_datos(self):
        texto = (self.carpeta / "entorno.txt").read_text(encoding="utf-8")
        self.assertIn("Semilla: 2026", texto)
        self.assertIn(f"Repeticiones: {self.REPETICIONES}", texto)
        self.assertIn(sys.version.split()[0], texto)
        self.assertIn("RAM: no disponible", texto)

    def test_ejecutar_dos_veces_sobrescribe_sin_error(self):
        with tempfile.TemporaryDirectory(
            prefix="tmp_experimento_", dir=CARPETA_TESTS
        ) as temporal:
            carpeta = Path(temporal) / "resultados"
            for _ in range(2):
                ExperimentoEficiencia(
                    repeticiones=1, carpeta_salida=carpeta
                ).ejecutar()
            self.assertEqual(len(leer_csv(carpeta / "arboles.csv")), 6)
            self.assertEqual(len(leer_csv(carpeta / "heap.csv")), 3)

    @unittest.skipUnless(find_spec("openpyxl"), "openpyxl no esta instalado")
    def test_excel_con_una_hoja_por_tabla(self):
        from openpyxl import load_workbook

        libro = load_workbook(self.carpeta / "resultados.xlsx")
        try:
            self.assertEqual(
                libro.sheetnames,
                ["arboles", "heap", "mediciones_crudas", "ordenes_usados"],
            )
            hoja = libro["arboles"]
            self.assertEqual(hoja.max_row, 7)
            self.assertIsInstance(hoja["D2"].value, int)
            self.assertIsInstance(hoja["F2"].value, float)
        finally:
            libro.close()

    @unittest.skipUnless(find_spec("matplotlib"), "matplotlib no esta instalado")
    def test_graficos_generados(self):
        for nombre in ("alturas.png", "tiempos.png"):
            ruta = self.carpeta / nombre
            self.assertTrue(ruta.is_file(), nombre)
            self.assertGreater(ruta.stat().st_size, 0)


if __name__ == "__main__":
    unittest.main()
